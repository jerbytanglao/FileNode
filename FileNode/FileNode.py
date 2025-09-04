import os
import threading
import time
import tempfile
from tkinter import filedialog
from customtkinter import (
    CTk, CTkFrame, CTkLabel, CTkButton, CTkProgressBar, CTkScrollableFrame, CTkEntry, CTkImage
)
from PIL import Image
from PyPDF2 import PdfMerger, PdfReader, PdfWriter
from docx2pdf import convert as docx2pdf_convert
from fpdf import FPDF
from docx import Document

# --- Logic Classes ---

class PDFTools:
    def merge_pdfs(self, file_list, output_path="merged_output.pdf"):
        merger = PdfMerger()
        for pdf in file_list:
            merger.append(pdf)
        merger.write(output_path)
        merger.close()
        return output_path

    def split_pdf(self, input_path, start_page, end_page, output_path="split_output.pdf"):
        reader = PdfReader(input_path)
        writer = PdfWriter()
        for i in range(start_page - 1, end_page):
            writer.add_page(reader.pages[i])
        with open(output_path, "wb") as f:
            writer.write(f)
        return output_path

    def convert_pdf_to_word(self, input_path, output_path="converted_output.docx"):
        # Simple text extraction (not formatting)
        reader = PdfReader(input_path)
        doc = Document()
        for page in reader.pages:
            text = page.extract_text()
            if text:
                doc.add_paragraph(text)
        doc.save(output_path)
        return output_path

class WordTools:
    def convert_docx_to_pdf(self, input_path, output_path="converted_output.pdf"):
        docx2pdf_convert(input_path, output_path)
        return output_path

class ImageTools:
    def reduce_image_size(self, input_path, output_path, quality=70):
        img = Image.open(input_path)
        img.save(output_path, quality=int(quality), optimize=True)
        return output_path

    def convert_image_format(self, input_path, output_path):
        img = Image.open(input_path)
        img.save(output_path)
        return output_path

class TextTools:
    def convert_text_to_pdf(self, input_path, output_path="converted_text.pdf"):
        with open(input_path, "r", encoding="utf-8") as f:
            lines = f.readlines()
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size=12)
        for line in lines:
            pdf.cell(200, 10, txt=line.strip(), ln=1)
        pdf.output(output_path)
        return output_path

# --- GUI ---

class MainWindow(CTk):
    def __init__(self):
        super().__init__()
        self.title("File Tool Application")
        self.geometry("1000x700")

        self.accent_color = "#1e90ff"
        self.button_bg = "#2a4e7c"
        self.button_hover = "#2776c7"
        self.universal_font = ("Segoe UI", 15, "bold")
        self.button_font = ("Segoe UI", 14, "bold")

        self.selected_tool = None
        self.uploaded_files = []
        self.progress_bar = None
        self.upload_button = None
        self.save_button = None

        self.create_menu()
        self.create_main_area()
        self.create_status_bar()

    def create_menu(self):
        self.menu_frame = CTkScrollableFrame(
            self, width=260, fg_color="#23272f", corner_radius=16
        )
        self.menu_frame.pack(side="left", fill="y", padx=12, pady=12)

        self.tool_buttons = {}

        tools = {
            "PDF Tools": ["Merge PDF", "Split PDF", "PDF to Word"],
            "Word Tools": ["Docs to PDF"],
            "Image Tools": ["Image Size Reducer", "Image Format Converter"],
            "Text Tools": ["Text to PDF"],
        }

        for section, tool_list in tools.items():
            section_label = CTkLabel(
                self.menu_frame,
                text=section,
                font=self.universal_font,
                text_color="#f5f6fa",
                pady=6,
            )
            section_label.pack(pady=(18, 6), anchor="center")
            for tool in tool_list:
                btn = CTkButton(
                    self.menu_frame,
                    text=tool,
                    height=36,
                    width=200,
                    font=self.button_font,
                    fg_color=self.button_bg,
                    hover_color=self.button_hover,
                    text_color="#f5f6fa",
                    corner_radius=10,
                    border_width=0,
                    command=lambda t=tool: self.select_tool(t)
                )
                btn.pack(pady=(0, 10), padx=8, anchor="center")
                self.tool_buttons[tool] = btn

    def create_main_area(self):
        self.workspace_frame = CTkFrame(self, fg_color="#1A1C23", corner_radius=22)
        self.workspace_frame.pack(side="left", fill="both", expand=True)

        self.thumbnail_panel = CTkFrame(self.workspace_frame, width=340, fg_color="#181A20", corner_radius=18)
        self.thumbnail_panel.pack(side="left", fill="both", expand=True, padx=16, pady=16)

        self.input_panel = CTkFrame(self.workspace_frame, fg_color="#232A34", corner_radius=18)
        self.input_panel.pack(side="right", fill="both", expand=True, padx=16, pady=16)

        self.info_label = CTkLabel(self.input_panel, text="Select a tool to continue", font=self.universal_font, text_color="#F7F8FA", fg_color="#232A34")
        self.info_label.pack(pady=(28, 18), anchor="n")

        self.button_row = CTkFrame(self.input_panel, fg_color="#232A34")
        self.button_row.pack(pady=(0, 28), anchor="n")

        self.upload_button = None
        self.cancel_button = CTkButton(self.button_row, text="Cancel", font=self.button_font, fg_color="#232A34", hover_color="#4F8CFF", text_color="#F7F8FA", corner_radius=16, command=self.cancel_operation)
        self.cancel_button.pack(side="left", padx=8)
        self.cancel_button.configure(state="disabled")

    def create_status_bar(self):
        self.status_label = CTkLabel(self, text="", font=self.universal_font)
        self.status_label.place(relx=1.0, rely=1.0, anchor="se", x=-10, y=-10)

    def select_tool(self, tool):
        self.selected_tool = tool
        self.uploaded_files = []
        self.info_label.configure(text=f"Selected Tool: {tool}")

        for t, btn in self.tool_buttons.items():
            if t == tool:
                btn.configure(fg_color=self.accent_color, hover_color=self.button_hover)
            else:
                btn.configure(fg_color=self.button_bg, hover_color=self.button_hover)

        if self.upload_button and self.upload_button.winfo_exists():
            self.upload_button.destroy()
        if self.save_button and self.save_button.winfo_exists():
            self.save_button.destroy()
        if self.progress_bar and self.progress_bar.winfo_exists():
            self.progress_bar.destroy()
        if hasattr(self, "reduce_button") and self.reduce_button and self.reduce_button.winfo_exists():
            self.reduce_button.destroy()
        if hasattr(self, "upload_again_button") and self.upload_again_button and self.upload_again_button.winfo_exists():
            self.upload_again_button.destroy()
        if hasattr(self, "size_info_label") and self.size_info_label and self.size_info_label.winfo_exists():
            self.size_info_label.destroy()

        for widget in self.input_panel.winfo_children():
            if widget not in [self.info_label, self.button_row]:
                widget.destroy()
        for widget in self.thumbnail_panel.winfo_children():
            widget.destroy()
        for widget in self.button_row.winfo_children():
            widget.destroy()

        if tool == "Merge PDF":
            self.upload_button = CTkButton(self.button_row, text="Upload PDF Files", font=self.button_font, command=self.upload_file_for_merge)
            self.upload_button.pack(side="left", padx=5)
            self.status_label.configure(text="Please upload 2 PDF files to merge.")
        elif tool == "Split PDF":
            self.upload_button = CTkButton(self.button_row, text="Upload PDF File", font=self.button_font, command=self.upload_file_for_split)
            self.upload_button.pack(side="left", padx=5)
            self.status_label.configure(text="Please upload a PDF file to split.")
        elif tool == "PDF to Word":
            self.upload_button = CTkButton(self.button_row, text="Upload PDF File", font=self.button_font, command=self.upload_file_for_pdf_to_word)
            self.upload_button.pack(side="left", padx=5)
            self.status_label.configure(text="Please upload a PDF file to convert to Word.")
        elif tool == "Docs to PDF":
            self.upload_button = CTkButton(self.button_row, text="Upload DOCX File", font=self.button_font, command=self.upload_file_for_docs_to_pdf)
            self.upload_button.pack(side="left", padx=5)
            self.status_label.configure(text="Please upload a DOCX file to convert to PDF.")
        elif tool == "Image Size Reducer":
            self.upload_button = CTkButton(self.button_row, text="Upload Image File", font=self.button_font, command=self.upload_file_for_image_size_reducer)
            self.upload_button.pack(side="left", padx=5)
            self.status_label.configure(text="Please upload an image file to reduce size.")
        elif tool == "Image Format Converter":
            self.upload_button = CTkButton(self.button_row, text="Upload Image File", font=self.button_font, command=self.upload_file_for_image_format_converter)
            self.upload_button.pack(side="left", padx=5)
            self.status_label.configure(text="Please upload an image file to convert format.")
        elif tool == "Text to PDF":
            self.upload_button = CTkButton(self.button_row, text="Upload Text File", font=self.button_font, command=self.upload_file_for_text_to_pdf)
            self.upload_button.pack(side="left", padx=5)
            self.status_label.configure(text="Please upload a text file to convert to PDF.")

        self.cancel_button = CTkButton(self.button_row, text="Cancel", font=self.button_font, command=self.cancel_operation)
        self.cancel_button.pack(side="left", padx=5)
        self.cancel_button.configure(state="disabled")

    # --- PDF Merge ---
    def upload_file_for_merge(self):
        file_paths = filedialog.askopenfilenames(filetypes=[("PDF Files", "*.pdf")])
        if file_paths and len(file_paths) == 2:
            self.uploaded_files = list(file_paths)
            self.info_label.configure(text="Files uploaded. Processing...")
            self.upload_button.configure(state="disabled")
            self.cancel_button.configure(state="normal")
            self.progress_bar = CTkProgressBar(self.input_panel, width=200)
            self.progress_bar.pack(pady=10)
            self.progress_bar.set(0)
            self.merge_ready = False
            threading.Thread(target=self.start_merge_process, daemon=True).start()
            threading.Thread(target=self.animate_progress, daemon=True).start()
        else:
            self.status_label.configure(text="Please select exactly 2 PDF files.")

    def start_merge_process(self):
        pdf_tools = PDFTools()
        output_path = pdf_tools.merge_pdfs(self.uploaded_files)
        self.merged_output_path = output_path
        self.merge_ready = True

    def animate_progress(self):
        val = 0
        while val < 1.0:
            val += 0.01
            self.progress_bar.set(val)
            time.sleep(0.03)
        self.progress_bar.set(1.0)
        while not getattr(self, "merge_ready", False):
            time.sleep(0.05)
        self.show_save_button(self.merged_output_path)
        self.status_label.configure(text="PDFs merged successfully!")
        self.info_label.configure(text="Merge complete. You can now save the file.")

    def show_save_button(self, output_path):
        if self.save_button and self.save_button.winfo_exists():
            self.save_button.destroy()
        def save_file():
            save_path = filedialog.asksaveasfilename(defaultextension=".pdf")
            if save_path:
                with open(output_path, "rb") as fsrc, open(save_path, "wb") as fdst:
                    fdst.write(fsrc.read())
                self.status_label.configure(text=f"File saved: {os.path.basename(save_path)}")
                self.info_label.configure(text="File saved successfully!")
        self.save_button = CTkButton(self.input_panel, text="Save PDF", font=self.button_font, command=save_file)
        self.save_button.pack(pady=10)

    # --- PDF Split ---
    def upload_file_for_split(self):
        file_path = filedialog.askopenfilename(filetypes=[("PDF Files", "*.pdf")])
        if file_path:
            self.uploaded_files = [file_path]
            self.info_label.configure(text="PDF uploaded. Enter the page range to split.")
            self.upload_button.configure(state="disabled")
            self.cancel_button.configure(state="normal")
            self.show_pdf_pages_left(file_path)
            self.show_split_pdf_ui_right(file_path)
        else:
            self.status_label.configure(text="Please select a PDF file.")

    def show_pdf_pages_left(self, pdf_path):
        for widget in self.thumbnail_panel.winfo_children():
            widget.destroy()
        try:
            from pdf2image import convert_from_path
            pages = convert_from_path(pdf_path, size=(120, 160))
            for idx, img in enumerate(pages, 1):
                pil_img = img
                pil_img.thumbnail((350, 500))
                ctk_img = CTkImage(light_image=pil_img, size=pil_img.size)
                lbl = CTkLabel(self.thumbnail_panel, image=ctk_img, text=f"Page {idx}", font=self.universal_font)
                lbl.image = ctk_img
                lbl.pack(pady=10)
        except Exception as e:
            CTkLabel(self.thumbnail_panel, text=f"Preview error: {e}", font=self.universal_font).pack(pady=5)

    def show_split_pdf_ui_right(self, file_path):
        for widget in self.input_panel.winfo_children():
            if widget not in [self.info_label, self.button_row]:
                widget.destroy()
        CTkLabel(self.input_panel, text=f"Selected PDF: {os.path.basename(file_path)}", font=self.universal_font).pack(pady=5)
        CTkLabel(self.input_panel, text="Enter page range (e.g., 1-5):", font=self.universal_font).pack(pady=5)
        self.page_range_entry = CTkEntry(self.input_panel, width=200, font=self.universal_font)
        self.page_range_entry.pack(pady=5)
        self.split_btn = CTkButton(self.input_panel, text="Split PDF", font=self.button_font, command=lambda: self.start_split_pdf(file_path))
        self.split_btn.pack(pady=10)
        self.split_save_btn = None

    def start_split_pdf(self, file_path):
        page_range = self.page_range_entry.get().strip()
        try:
            start_page, end_page = map(int, page_range.split("-"))
            if start_page > end_page:
                raise ValueError("Start page must be less than or equal to end page.")
        except Exception:
            self.status_label.configure(text="❌ Invalid page range. Use the format: start-end (e.g., 1-5).")
            return
        self.status_label.configure(text="Splitting PDF, please wait...")
        self.progress_bar = CTkProgressBar(self.input_panel, width=200)
        self.progress_bar.pack(pady=10)
        self.progress_bar.set(0)
        if hasattr(self, "split_save_btn") and self.split_save_btn and self.split_save_btn.winfo_exists():
            self.split_save_btn.destroy()
        self.split_ready = False
        threading.Thread(target=self.split_pdf_process, args=(file_path, start_page, end_page), daemon=True).start()
        threading.Thread(target=self.animate_split_progress, daemon=True).start()

    def animate_split_progress(self):
        val = 0
        while val < 1.0:
            val += 0.01
            self.progress_bar.set(val)
            time.sleep(0.03)
        self.progress_bar.set(1.0)
        while not getattr(self, "split_ready", False):
            time.sleep(0.05)
        self.show_save_button_split(self.split_output_path)

    def split_pdf_process(self, file_path, start_page, end_page):
        pdf_tools = PDFTools()
        output_path = "split_output.pdf"
        try:
            pdf_tools.split_pdf(file_path, start_page, end_page, output_path)
            self.split_output_path = output_path
            self.split_ready = True
            self.status_label.configure(text="PDF split successfully!")
            self.info_label.configure(text="Split complete. You can now save the file.")
        except Exception as e:
            self.status_label.configure(text=f"❌ Error: {e}")

    def show_save_button_split(self, output_path):
        if hasattr(self, "split_save_btn") and self.split_save_btn and self.split_save_btn.winfo_exists():
            self.split_save_btn.destroy()
        def save_file():
            save_path = filedialog.asksaveasfilename(defaultextension=".pdf")
            if save_path:
                with open(output_path, "rb") as fsrc, open(save_path, "wb") as fdst:
                    fdst.write(fsrc.read())
                self.status_label.configure(text=f"File saved: {os.path.basename(save_path)}")
                self.info_label.configure(text="File saved successfully!")
        self.split_save_btn = CTkButton(self.input_panel, text="Save Split PDF", font=self.button_font, command=save_file)
        self.split_save_btn.pack(pady=10)

    # --- PDF to Word ---
    def upload_file_for_pdf_to_word(self):
        file_path = filedialog.askopenfilename(filetypes=[("PDF Files", "*.pdf")])
        if file_path:
            self.uploaded_files = [file_path]
            self.info_label.configure(text="PDF uploaded. Converting to Word...")
            self.upload_button.configure(state="disabled")
            self.cancel_button.configure(state="normal")
            self.show_pdf_pages_left(file_path)
            self.progress_bar = CTkProgressBar(self.input_panel, width=200)
            self.progress_bar.pack(pady=10)
            self.progress_bar.set(0)
            self.pdf_to_word_ready = False
            threading.Thread(target=self.start_pdf_to_word_process, args=(file_path,), daemon=True).start()
            threading.Thread(target=self.animate_pdf_to_word_progress, daemon=True).start()
        else:
            self.status_label.configure(text="Please select a PDF file.")

    def start_pdf_to_word_process(self, file_path):
        pdf_tools = PDFTools()
        output_path = "converted_output.docx"
        try:
            pdf_tools.convert_pdf_to_word(file_path, output_path)
            self.converted_word_path = output_path
            self.pdf_to_word_ready = True
            self.status_label.configure(text="PDF converted successfully!")
        except Exception as e:
            self.status_label.configure(text=f"❌ Error: {e}")

    def animate_pdf_to_word_progress(self):
        val = 0
        while val < 1.0:
            val += 0.01
            self.progress_bar.set(val)
            time.sleep(0.03)
        self.progress_bar.set(1.0)
        while not getattr(self, "pdf_to_word_ready", False):
            time.sleep(0.05)
        self.info_label.configure(text="Conversion complete. You can now save the Word file.")
        self.show_save_button_word(self.converted_word_path)

    def show_save_button_word(self, output_path):
        if hasattr(self, "save_button") and self.save_button and self.save_button.winfo_exists():
            self.save_button.destroy()
        def save_file():
            save_path = filedialog.asksaveasfilename(defaultextension=".docx")
            if save_path:
                with open(output_path, "rb") as fsrc, open(save_path, "wb") as fd:
                    fd.write(fsrc.read())
                self.status_label.configure(text=f"File saved: {os.path.basename(save_path)}")
                self.info_label.configure(text="File saved successfully!")
        self.save_button = CTkButton(self.input_panel, text="Save Word File", font=self.button_font, command=save_file)
        self.save_button.pack(pady=10)

    # --- Docs to PDF ---
    def upload_file_for_docs_to_pdf(self):
        file_path = filedialog.askopenfilename(filetypes=[("Word Documents", "*.docx")])
        if file_path:
            self.uploaded_files = [file_path]
            self.info_label.configure(text="DOCX uploaded. Converting to PDF...")
            self.upload_button.configure(state="disabled")
            self.cancel_button.configure(state="normal")
            self.progress_bar = CTkProgressBar(self.input_panel, width=200)
            self.progress_bar.pack(pady=10)
            self.progress_bar.set(0)
            self.docs_to_pdf_ready = False
            threading.Thread(target=self.start_docs_to_pdf_process, args=(file_path,), daemon=True).start()
            threading.Thread(target=self.animate_docs_to_pdf_progress, daemon=True).start()
        else:
            self.status_label.configure(text="Please select a DOCX file.")

    def start_docs_to_pdf_process(self, file_path):
        word_tools = WordTools()
        output_path = "converted_output.pdf"
        try:
            word_tools.convert_docx_to_pdf(file_path, output_path)
            self.converted_pdf_path = output_path
            self.docs_to_pdf_ready = True
            self.status_label.configure(text="DOCX converted successfully!")
        except Exception as e:
            self.status_label.configure(text=f"❌ Error: {e}")

    def animate_docs_to_pdf_progress(self):
        val = 0
        while val < 1.0:
            val += 0.01
            self.progress_bar.set(val)
            time.sleep(0.03)
        self.progress_bar.set(1.0)
        while not getattr(self, "docs_to_pdf_ready", False):
            time.sleep(0.05)
        self.info_label.configure(text="Conversion complete. You can now save the PDF file.")
        self.show_save_button(self.converted_pdf_path)

    # --- Image Size Reducer ---
    def format_size(self, size_bytes):
        if size_bytes < 1000:
            return f"{size_bytes} B"
        elif size_bytes < 1_000_000:
            return f"{size_bytes/1000:.2f} KB"
        elif size_bytes < 1_000_000_000:
            return f"{size_bytes/1_000_000:.2f} MB"
        else:
            return f"{size_bytes/1_000_000_000:.2f} GB"

    def upload_file_for_image_size_reducer(self):
        file_path = filedialog.askopenfilename(filetypes=[("Image Files", "*.png;*.jpg;*.jpeg;*.bmp")])
        if file_path:
            self.uploaded_files = [file_path]
            file_size = os.path.getsize(file_path)
            formatted_size = self.format_size(file_size)
            self.original_image_size = file_size
            self.current_image_path = file_path
            self.info_label.configure(
                text=f"Image uploaded: {os.path.basename(file_path)}"
            )
            self.upload_button.configure(state="disabled")
            self.cancel_button.configure(state="normal")
            # Show thumbnail
            for widget in self.thumbnail_panel.winfo_children():
                widget.destroy()
            pil_img = Image.open(file_path)
            pil_img.thumbnail((320, 320))
            ctk_img = CTkImage(light_image=pil_img, size=pil_img.size)
            lbl = CTkLabel(self.thumbnail_panel, image=ctk_img, text="")
            lbl.image = ctk_img
            lbl.pack(pady=10)
            # Remove previous reduce button if exists
            if hasattr(self, "reduce_button") and self.reduce_button and self.reduce_button.winfo_exists():
                self.reduce_button.destroy()
            # Remove previous save button if exists
            if hasattr(self, "save_button") and self.save_button and self.save_button.winfo_exists():
                self.save_button.destroy()
            # Remove previous size info if exists
            if hasattr(self, "size_info_label") and self.size_info_label and self.size_info_label.winfo_exists():
                self.size_info_label.destroy()
            # Show original size info below buttons
            self.size_info_label = CTkLabel(self.input_panel, text=f"Original size: {formatted_size}", font=self.universal_font, text_color="#F7F8FA")
            self.size_info_label.pack(pady=(10, 0))
            # Add Reduce button
            self.reduce_button = CTkButton(self.input_panel, text="Reduce", font=self.button_font, command=lambda: self.start_image_size_reduce_process(file_path, file_size))
            self.reduce_button.pack(pady=10)
        else:
            self.status_label.configure(text="Please select an image file.")

    def start_image_size_reduce_process(self, file_path, original_size):
        self.progress_bar = CTkProgressBar(self.input_panel, width=200)
        self.progress_bar.pack(pady=10)
        self.progress_bar.set(0)
        self.image_reduce_ready = False
        self.original_image_size = original_size
        self.current_image_path = file_path
        # Remove previous save button if exists
        if hasattr(self, "save_button") and self.save_button and self.save_button.winfo_exists():
            self.save_button.destroy()
        # Remove reduce button after click
        if hasattr(self, "reduce_button") and self.reduce_button and self.reduce_button.winfo_exists():
            self.reduce_button.destroy()
        threading.Thread(target=self._reduce_image_thread, args=(file_path,), daemon=True).start()
        threading.Thread(target=self.animate_image_reduce_progress, daemon=True).start()

    def _reduce_image_thread(self, file_path):
        image_tools = ImageTools()
        output_path = "reduced_image.jpg"
        try:
            image_tools.reduce_image_size(file_path, output_path)
            self.reduced_image_path = output_path
            self.image_reduce_ready = True
            self.status_label.configure(text="Image size reduced successfully!")
        except Exception as e:
            self.status_label.configure(text=f"❌ Error: {e}")

    def animate_image_reduce_progress(self):
        val = 0
        while val < 1.0:
            val += 0.01
            self.progress_bar.set(val)
            time.sleep(0.03)
        self.progress_bar.set(1.0)
        while not getattr(self, "image_reduce_ready", False):
            time.sleep(0.05)
        # Show reduced file size, original file size, and percentage reduced below the buttons
        reduced_size = os.path.getsize(self.reduced_image_path)
        original_size = getattr(self, "original_image_size", 0)
        percent_reduced = 0
        if original_size > 0:
            percent_reduced = 100 - (reduced_size / original_size * 100)
        # Remove previous size info if exists
        if hasattr(self, "size_info_label") and self.size_info_label and self.size_info_label.winfo_exists():
            self.size_info_label.destroy()
        self.size_info_label = CTkLabel(
            self.input_panel,
            text=(
                f"Original size: {self.format_size(original_size)}\n"
                f"Reduced size: {self.format_size(reduced_size)}\n"
                f"Size reduced: {percent_reduced:.1f}%"
            ),
            font=self.universal_font,
            text_color="#F7F8FA"
        )
        self.size_info_label.pack(pady=(10, 0))
        self.show_save_button_image(self.reduced_image_path)
        self.show_upload_again_button()

    def show_save_button_image(self, output_path):
        if hasattr(self, "save_button") and self.save_button and self.save_button.winfo_exists():
            self.save_button.destroy()
        def save_file():
            save_path = filedialog.asksaveasfilename(defaultextension=".jpg")
            if save_path:
                with open(output_path, "rb") as fsrc, open(save_path, "wb") as fdst:
                    fdst.write(fsrc.read())
                self.status_label.configure(text=f"File saved: {os.path.basename(save_path)}")
                self.info_label.configure(text="File saved successfully!")
        self.save_button = CTkButton(self.input_panel, text="Save Image", font=self.button_font, command=save_file)
        self.save_button.pack(pady=10)

    def show_upload_again_button(self):
        if hasattr(self, "upload_again_button") and self.upload_again_button and self.upload_again_button.winfo_exists():
            self.upload_again_button.destroy()
        self.upload_again_button = CTkButton(
            self.input_panel,
            text="Upload Image File",
            font=self.button_font,
            command=self.upload_file_for_image_size_reducer
        )
        self.upload_again_button.pack(pady=10)

    # --- Image Format Converter ---
    def upload_file_for_image_format_converter(self):
        file_path = filedialog.askopenfilename(filetypes=[("Image Files", "*.png;*.jpg;*.jpeg;*.bmp")])
        if file_path:
            self.uploaded_files = [file_path]
            self.info_label.configure(text="Image uploaded. Converting format...")
            self.upload_button.configure(state="disabled")
            self.cancel_button.configure(state="normal")
            self.progress_bar = CTkProgressBar(self.input_panel, width=200)
            self.progress_bar.pack(pady=10)
            self.progress_bar.set(0)
            self.image_convert_ready = False
            threading.Thread(target=self.start_image_format_convert_process, args=(file_path,), daemon=True).start()
            threading.Thread(target=self.animate_image_convert_progress, daemon=True).start()
        else:
            self.status_label.configure(text="Please select an image file.")

    def start_image_format_convert_process(self, file_path):
        image_tools = ImageTools()
        output_path = "converted_image.png"
        try:
            image_tools.convert_image_format(file_path, output_path)
            self.converted_image_path = output_path
            self.image_convert_ready = True
            self.status_label.configure(text="Image format converted successfully!")
        except Exception as e:
            self.status_label.configure(text=f"❌ Error: {e}")

    def animate_image_convert_progress(self):
        val = 0
        while val < 1.0:
            val += 0.01
            self.progress_bar.set(val)
            time.sleep(0.03)
        self.progress_bar.set(1.0)
        while not getattr(self, "image_convert_ready", False):
            time.sleep(0.05)
        self.info_label.configure(text="Conversion complete. You can now save the image.")
        self.show_save_button_image(self.converted_image_path)
        self.show_upload_again_button()

    # --- Text to PDF ---
    def upload_file_for_text_to_pdf(self):
        file_path = filedialog.askopenfilename(filetypes=[("Text Files", "*.txt")])
        if file_path:
            self.uploaded_files = [file_path]
            self.info_label.configure(text="Text file uploaded. Converting to PDF...")
            self.upload_button.configure(state="disabled")
            self.cancel_button.configure(state="normal")
            self.progress_bar = CTkProgressBar(self.input_panel, width=200)
            self.progress_bar.pack(pady=10)
            self.progress_bar.set(0)
            self.text_to_pdf_ready = False
            threading.Thread(target=self.start_text_to_pdf_process, args=(file_path,), daemon=True).start()
            threading.Thread(target=self.animate_text_to_pdf_progress, daemon=True).start()
        else:
            self.status_label.configure(text="Please select a text file.")

    def start_text_to_pdf_process(self, file_path):
        text_tools = TextTools()
        output_path = "converted_text.pdf"
        try:
            text_tools.convert_text_to_pdf(file_path, output_path)
            self.converted_text_pdf_path = output_path
            self.text_to_pdf_ready = True
            self.status_label.configure(text="Text converted successfully!")
        except Exception as e:
            self.status_label.configure(text=f"❌ Error: {e}")

    def animate_text_to_pdf_progress(self):
        val = 0
        while val < 1.0:
            val += 0.01
            self.progress_bar.set(val)
            time.sleep(0.03)
        self.progress_bar.set(1.0)
        while not getattr(self, "text_to_pdf_ready", False):
            time.sleep(0.05)
        self.info_label.configure(text="Conversion complete. You can now save the PDF file.")
        self.show_save_button(self.converted_text_pdf_path)

    # --- Cancel ---
    def cancel_operation(self):
        for panel in [self.thumbnail_panel, self.input_panel]:
            for widget in panel.winfo_children():
                widget.destroy()
        self.info_label = CTkLabel(self.input_panel, text="Select a tool to continue", font=self.universal_font, text_color="#F7F8FA", fg_color="#232A34")
        self.info_label.pack(pady=(28, 18), anchor="n")
        self.button_row = CTkFrame(self.input_panel, fg_color="#232A34")
        self.button_row.pack(pady=(0, 28), anchor="n")
        self.cancel_button = CTkButton(self.button_row, text="Cancel", font=self.button_font, fg_color="#232A34", hover_color="#4F8CFF", text_color="#F7F8FA", corner_radius=16, command=self.cancel_operation)
        self.cancel_button.pack(side="left", padx=8)
        self.cancel_button.configure(state="disabled")
        self.upload_button = None
        self.selected_tool = None
        self.status_label.configure(text="")

if __name__ == "__main__":
    app = MainWindow()
    app.mainloop()